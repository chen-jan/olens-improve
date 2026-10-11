"""Invented full-menu rates exercise source guards and editable geometry only."""
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import build_expansion_slides as b


def write(path,value):
    path.write_text(json.dumps(value,indent=2)+'\n')
    return b.digest(path)


def fixture(root,qa=True):
    """No measured results are read; only the frozen scientific menu is reused."""
    inventory=json.loads((ROOT/'configs/expansion_manifest_v1.json').read_text())
    if qa:inventory['synthetic_layout_fixture']=True
    ip=root/'inventory.json';isha=write(ip,inventory)
    active=[r for r in inventory['conditions'] if r['status'] in b.ACTIVE]
    rows=[];lengths={};secondary={}
    for i,c in enumerate(active):
        cid=c['condition_id'];lengths[cid]={};secondary[cid]={}
        for j,f in enumerate(b.FAMILIES):
            den=100;num=15+(i*3+j*7)%70;value=num/den
            rows.append({'condition':cid,'family':f,'metric':{
                'name':'original_grid_contradiction_rate' if j==0 else 'original_grid_item_pass_rate',
                'value':value,'numerator':num,'denominator':den,
                'ci95':[max(0,value-.09),min(1,value+.09)],'items':[149,100,120,100][j],
                'independent_groups':inventory['family_group_counts'][f],'defined_draws':10000}})
            lengths[cid][f]={'mean_qwen_tokens_all_intended':320 if cid=='raw_n1-n1' else 45+(i*7)%175}
            secondary[cid][f]={'coverage':{'empty_abstentions':i%3,'expected_cells':100},
                              'assertion_share':{'numerator':80,'denominator':100},
                              'cell_recovery':{'numerator':70,'denominator':100}}
    lookup={(r['condition'],r['family']):r['metric'] for r in rows};comparisons=[]
    for i,c in enumerate(inventory['headline_comparison_family']):
        for f in b.FAMILIES:
            delta=lookup[c['condition'],f]['value']-lookup[c['reference'],f]['value']
            comparisons.append({**c,'family':f,'delta':delta,'ci95_delta':[delta-.08,delta+.08],
                                'simultaneous_ci95_delta':None if i==0 else [delta-.21,delta+.21],
                                'bootstrap_standard_error':0 if i==0 else .04,'defined_draws':10000})
    # Undefined ratios must stay visible, while exact counts remain intact.
    lookup[active[-1]['condition_id'],b.FAMILIES[0]].update(value=None,numerator=0,denominator=0,ci95=None)
    results={'status':'synthetic_layout_fixture' if qa else 'complete_exploratory','manifest_sha256':isha,
             'missing_condition_families':[],'cohorts':{'full':{'rows':rows,'comparisons':comparisons,
             'comparison_family_complete':True,'missing_comparisons':[],'lengths':lengths,
             'secondary':secondary,'bootstrap_replicates':10000,'declared_contrasts':len(comparisons)}}}
    if qa:results['synthetic_layout_fixture']=True
    rp=root/'results.json';rsha=write(rp,results)
    examples=[]
    for i,f in enumerate((b.FAMILIES[0],b.FAMILIES[2])):
        examples.append({'condition':'best_n_exact-n32','family':f,'cell_id':f+'|synthetic|L12|p7',
                         'reference':'capped_n1-n1','baseline_text':'- Synthetic original description.\n- A retained wrong statement.',
                         'selected_text':'- Synthetic selected description.\n- A retained right statement.',
                         'interpretation':'Invented text illustrates layout only; it supplies no scientific finding.',
                         **({'synthetic_layout_fixture':True} if qa else {})})
    if not qa:
        pack={'manifest_sha256':isha,'contexts':{},'conditions':{'best_n_exact-n32':{'families':{}}}}
        for e in examples:pack['conditions'][e['condition']]['families'][e['family']]={
            'fixed_anchor':[{'cell_id':e['cell_id'],'baseline_condition':e['reference'],
                             'baseline':{'text':e['baseline_text']},'proposed':{'text':e['selected_text']}}]}
        pp=root/'pack.json';psha=write(pp,pack)
        for e in examples:e.update(source_path=str(pp),source_sha256=psha)
    review={'status':'synthetic_layout_review' if qa else 'root_reviewed','results_sha256':rsha,'manifest_sha256':isha,
            'title':'Expanded comparison · synthetic layout QA',
            'headline':'Invented fixture values test layout. They are not research results.',
            'findings':[{'title':'Synthetic finding placeholder','detail':'All chart rates are invented for rendering tests. No winner or recommendation is inferred.'}],
            'method_findings':{g:{'finding':'Synthetic layout coverage for this approach.','limits':'Not empirical evidence.'} for g in b.GROUPS},
            'recommendations':['Do not use synthetic slides as scientific evidence.'],
            'resource_summary':'Fixture only: no API request, GPU job, new statistic or measured cost was generated.',
            'presentation':{'examples':examples}}
    approaches={}
    for group in b.GROUPS:
        members=[r['condition_id'] for r in active if b.category(r['method'])==group]
        pair=next(c for c in inventory['headline_comparison_family'] if c['condition'] in members)
        e=deepcopy(examples[0]);e['condition']=pair['condition']
        approaches[group]={'action_title':b.GROUPS[group]+' · synthetic evidence only',
                           'bullets':['Synthetic quantitative layout; no actual finding.','Exact variants remain in the full appendix.'],
                           'mechanism_hypothesis':'Fixture hypothesis only: improving fit need not improve truthful content.',
                           'example':e,'contrast':{'condition':pair['condition'],'reference':pair['reference']}}
        if not qa:
            pack['conditions'].setdefault(e['condition'],{'families':{}})['families'][e['family']]={
                'fixed_anchor':[{'cell_id':e['cell_id'],'baseline_condition':e['reference'],
                                 'baseline':{'text':e['baseline_text']},'proposed':{'text':e['selected_text']}}]}
    if not qa:
        psha=write(pp,pack)
        for e in examples+[a['example'] for a in approaches.values()]:e['source_sha256']=psha
    review['presentation']['approaches']=approaches
    if qa:review['synthetic_layout_fixture']=True
    vp=root/'review.json';write(vp,review)
    return (rp,ip,vp),results,inventory,review


class ExpansionSlideContract(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)

    def test_full_menu_geometry_and_editable_output(self):
        paths,_,inventory,_=fixture(self.root)
        inputs=b.load(*paths,qa=True);pages,refs,core=b.pages_from(inputs,qa=True)
        self.assertEqual(core,25)
        self.assertEqual(len(refs),36)
        self.assertEqual(len(pages),25+14+31+1)
        self.assertEqual(set(inputs[3]),set(r['condition_id'] for r in inventory['conditions'] if r['status'] in b.ACTIVE))
        self.assertEqual(len(inputs[8]),496)
        self.assertIn('zero bootstrap variation',' '.join(p['notes'] for p in pages))
        self.assertIn('Undefined',' '.join(i.get('text','') for p in pages for i in p['items']))
        for p in pages:
            self.assertIn('SYNTHETIC LAYOUT QA',p['items'][1]['text'])
            for i in p['items']:
                self.assertIn(i['kind'],('text','rect','line','ellipse'))
                self.assertTrue(0<=i['x']<=b.W and 0<=i['y']<=b.H)
                if i['kind']!='line':self.assertLessEqual(i['x']+i['w'],b.W);self.assertLessEqual(i['y']+i['h'],b.H)
        out=self.root/'output';receipt=b.build(*paths,out,qa=True)
        from pptx import Presentation
        deck=Presentation(out/'expanded.pptx')
        self.assertEqual(len(deck.slides),71)
        self.assertEqual(receipt['inference_calls'],0)
        self.assertTrue(all(shape.shape_type!=13 for s in deck.slides for shape in s.shapes))
        self.assertIn('Synthetic original description',deck.slides[5].notes_slide.notes_text_frame.text)
        with self.assertRaisesRegex(ValueError,'immutable'):b.build(*paths,out,qa=True)

    def test_synthetic_inputs_rejected_by_final_mode(self):
        paths,*_=fixture(self.root)
        with self.assertRaisesRegex(ValueError,'Synthetic QA'):b.load(*paths)

    def test_final_mode_exact_review_source_and_complete_rows(self):
        paths,results,_,review=fixture(self.root,qa=False)
        self.assertEqual(len(b.load(*paths)[4]),280)
        review['results_sha256']='0'*64;write(paths[2],review)
        with self.assertRaisesRegex(ValueError,'reviewed'):b.load(*paths)
        results['cohorts']['full']['rows'].pop();review['results_sha256']=write(paths[0],results);write(paths[2],review)
        with self.assertRaisesRegex(ValueError,'every active'):b.load(*paths)

    def test_final_case_text_and_reference_are_exact(self):
        paths,_,_,review=fixture(self.root,qa=False)
        review['presentation']['examples'][0]['baseline_text']='Fabricated replacement'
        write(paths[2],review)
        with self.assertRaisesRegex(ValueError,'exact prose'):b.load(*paths)
        review['presentation']['examples'][0]['baseline_text']='- Synthetic original description.\n- A retained wrong statement.'
        review['presentation']['examples'][0]['reference']='raw_n1-n1';write(paths[2],review)
        with self.assertRaisesRegex(ValueError,'exact prose'):b.load(*paths)

    def test_frozen_contrast_completion_required(self):
        paths,results,_,review=fixture(self.root)
        results['cohorts']['full']['comparisons'].pop();review['results_sha256']=write(paths[0],results);write(paths[2],review)
        with self.assertRaisesRegex(ValueError,'frozen|Frozen'):b.load(*paths,qa=True)

    def parent_fixture(self):
        paths,results,inventory,review=fixture(self.root,qa=False)
        source=Path(review['presentation']['examples'][0]['source_path'])
        pack=json.loads(source.read_text())
        for group in ('pruning','combinations','editing'):
            a=review['presentation']['approaches'][group];e=a['example']
            condition=next(c for c in inventory['conditions'] if c['condition_id']==e['condition'])
            reference=next(p['reference'] for p in inventory['headline_comparison_family']
                           if p['condition']==e['condition'] and p['reference'] in condition['parent_condition_ids']
                           and p['reference'] not in ('raw_n1-n1','capped_n1-n1'))
            row=pack['conditions'][e['condition']]['families'][e['family']]['fixed_anchor'][0]
            row['original_baseline']=row.pop('baseline');row['original_baseline_condition']=row.pop('baseline_condition')
            row['parents']={reference:{'status':'complete','output':{'text':'- Synthetic immediate parent.'},
                                      'endpoint_evidence':{'synthetic':True}}}
            pack['conditions'][e['condition']]['frozen_parent_condition_ids']=condition['parent_condition_ids']
            e['reference']=reference;e['baseline_text']=row['parents'][reference]['output']['text']
            a['contrast']['reference']=reference
        psha=write(source,pack)
        for e in review['presentation']['examples']+[a['example'] for a in review['presentation']['approaches'].values()]:
            e['source_sha256']=psha
        write(paths[2],review)
        return paths,review,source,pack

    def test_pruning_combination_and_editing_parent_pack_examples(self):
        paths,review,_,_=self.parent_fixture()
        inputs=b.load(*paths)
        for group in ('pruning','combinations','editing'):
            e=inputs[6]['approaches'][group]['example']
            self.assertEqual(e['_reference_role'],'parent')
            self.assertEqual(e['_source_reference']['text'],'- Synthetic immediate parent.')
            self.assertIn('Synthetic original',e['_original_baseline']['text'])
        pages,_,_=b.pages_from(inputs)
        self.assertGreaterEqual(sum('Parent example' in i.get('text','') for p in pages for i in p['items']),3)

    def test_parent_pack_keeps_original_baseline_resolution(self):
        paths,review,_,_=self.parent_fixture()
        e=review['presentation']['approaches']['pruning']['example']
        e['reference']='capped_n1-n1';e['baseline_text']='- Synthetic original description.\n- A retained wrong statement.'
        write(paths[2],review)
        self.assertEqual(b.load(*paths)[6]['approaches']['pruning']['example']['_reference_role'],'original_baseline')

    def test_parent_resolution_rejects_missing_undeclared_or_wrong_text(self):
        paths,review,source,pack=self.parent_fixture()
        e=review['presentation']['approaches']['pruning']['example']
        condition=pack['conditions'][e['condition']];row=condition['families'][e['family']]['fixed_anchor'][0]
        for change in ('unavailable','wrong_text','undeclared','wrong_reference'):
            with self.subTest(change=change):
                altered=deepcopy(pack);c=altered['conditions'][e['condition']];r=c['families'][e['family']]['fixed_anchor'][0]
                if change=='unavailable':r['parents'][e['reference']]['status']='missing'
                elif change=='wrong_text':r['parents'][e['reference']]['output']['text']='Different actual parent'
                elif change=='undeclared':c['frozen_parent_condition_ids']=[]
                else:r['parents']={'best_n_exact-n16':r['parents'][e['reference']]}
                with self.assertRaisesRegex(ValueError,'exact prose'):b.resolve_example(altered,e)

    def test_reconstruction_is_separate_and_source_bound(self):
        paths,_,_,review=fixture(self.root)
        x={'manifest_sha256':review['manifest_sha256'],'quality_analysis_sha256':review['results_sha256'],
           'objective':'1-squared_residual/squared_target','solver':'scipy_exact',
           'vector_contract_sha256':'synthetic','rows':[]}
        for a in review['presentation']['approaches'].values():
            pair=a['contrast']
            for f in b.FAMILIES:x['rows'].append({'condition':pair['condition'],'family':f,
                'full_site_count':100,'scored_site_count':100,'unavailable_reasons':[],
                'mean_fve':.42,'median_fve':.41,'parent_comparisons':[
                    {'reference':pair['reference'],'paired_sites':100,'mean_delta':.025,
                     'ci95_delta':[.01,.04],'independent_groups':40,'warning':'Synthetic fixture only'}]})
        xp=self.root/'fit.json';review['presentation']['reconstruction_sha256']=write(xp,x);write(paths[2],review)
        with self.assertRaisesRegex(ValueError,'must be supplied'):b.load(*paths,qa=True)
        inputs=b.load(*paths,qa=True,reconstruction_path=xp);pages,_,core=b.pages_from(inputs,qa=True)
        self.assertEqual(core,26)
        self.assertIn('Reconstruction fit across the sample budgets',[p['title'] for p in pages])
        alltext=' '.join(i.get('text','') for p in pages for i in p['items'])
        self.assertIn('FVE Δ (pp) vs',alltext)
        self.assertIn('+2.5 pp',alltext)
        x['quality_analysis_sha256']='other';review['presentation']['reconstruction_sha256']=write(xp,x);write(paths[2],review)
        with self.assertRaisesRegex(ValueError,'bound'):b.load(*paths,qa=True,reconstruction_path=xp)

    def test_numeric_parent_can_differ_from_quality_reference(self):
        row={'parent_comparisons':[
            {'reference':'best_n_exact-n32','kind':'frozen_parent','mean_delta':-.003,
             'ci95_delta':[-.005,-.001],'paired_sites':98},
            {'reference':'raw_n1-n1','kind':'sampling_prefix_diagnostic','mean_delta':.02}]}
        actual=b.reconstruction_comparison(row,'capped_n1-n1')
        self.assertEqual(actual['reference'],'best_n_exact-n32')
        self.assertEqual(actual['paired_sites'],98)
        self.assertEqual(b.reconstruction_comparison(row,'raw_n1-n1')['reference'],'raw_n1-n1')
        self.assertIsNone(b.reconstruction_comparison({},'capped_n1-n1'))

    def test_reconstruction_scaling_percent_scope_and_actual_parent(self):
        rows=[]
        for f in b.FAMILIES:
            for n in b.N_VALUES:
                rows.append({'condition':f'best_n_exact-n{n}','family':f,
                    'mean_fve':.2+n/1000,'mean_fve_ci95':[.18,.26],
                    'status':'complete_comparable','mean_fve_warnings':[],
                    'scored_site_count':100,'full_site_count':100,
                    'parent_comparisons':[{'reference':'capped_n1-n1',
                        'paired_sites':97,'mean_delta':.025,'ci95_delta':[.01,.04]}]})
            for n in (16,32):
                rows.append({'condition':f'pooled_greedy_swap-n{n}','family':f,
                    'mean_fve':.3,'mean_fve_ci95':[.28,.32],
                    'status':'complete_comparable','parent_comparisons':[
                        {'reference':'best_n_exact-n32','paired_sites':96,
                         'mean_delta':.065,'ci95_delta':[.04,.08]}]})
            rows.append({'condition':'capped_n1-n1','family':f,
                'mean_fve':None,'mean_fve_ci95':None,'status':'partial_or_unavailable',
                'scored_site_count':0,'full_site_count':100,'parent_comparisons':[]})
            # This raw fit must never be silently substituted for the capped row.
            rows.append({'condition':'raw_n1-n1','family':f,'mean_fve':.99})
        items,notes=b.reconstruction_scaling(rows,{})
        visible=' '.join(i.get('text','') for i in items)
        self.assertIn('50%',visible)
        self.assertIn('Best32 vs capped: +2.5 pp',visible)
        self.assertIn('Pool32 vs Best32: +6.5 pp',visible)
        self.assertIn('†',visible)
        self.assertEqual(sum(i['kind']=='rect' and i['color']=='B57123' for i in items),8)
        self.assertFalse(any(i['kind']=='line' and i['color']==b.TEAL for i in items))
        self.assertIn('"paired_sites": 97',notes)
        self.assertNotIn('"condition": "raw_n1-n1"',notes)
        # Negative recorded fit/CI must extend the axis, without clipping.
        rows[0].update(mean_fve=-.2,mean_fve_ci95=[-.25,-.1])
        items,_=b.reconstruction_scaling(rows,{})
        self.assertIn('-30%', ' '.join(i.get('text','') for i in items))

    def test_final_build_requires_numeric_artifact(self):
        with self.assertRaisesRegex(ValueError,'source-bound reconstruction'):
            b.build(None,None,None,self.root/'final')

    def test_native_task_labels_and_boundary_warning_are_visible(self):
        paths,results,_,review=fixture(self.root)
        metric=results['cohorts']['full']['rows'][0]['metric']
        metric.update(value=0,numerator=0,ci95=[0,0],warnings=['Zero bootstrap variation'])
        review['results_sha256']=write(paths[0],results);write(paths[2],review)
        pages,_,_=b.pages_from(b.load(*paths,qa=True),qa=True)
        alltext=' '.join(i.get('text','') for p in pages for i in p['items'])
        self.assertIn('Hidden intermediate names',alltext)
        self.assertIn('Required intermediate entities; final answer is optional.',alltext)
        self.assertIn('0.0%†',alltext)
        self.assertIn('Zero bootstrap variation',' '.join(p['notes'] for p in pages))


if __name__=='__main__':unittest.main()
