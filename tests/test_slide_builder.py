"""Draft slide source binding and editable-output contract; no inference/API."""
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

spec = importlib.util.spec_from_file_location('slide_builder', Path(__file__).parents[1]/'scripts/build_slides.py')
builder = importlib.util.module_from_spec(spec); spec.loader.exec_module(builder)


def fixture(root):
    metrics = [{'family': f, 'condition': c, 'metric': {'value': .12, 'ci95': [.08, .2], 'numerator': 12, 'denominator': 100}}
               for f in builder.FAMILIES for c in builder.METHODS]
    section = {'stage': 'validation', 'review': {'status': 'root_reviewed'}, 'source_sha256': 'a'*64, 'cell_rows': metrics,
               'family_counts': {f: {'items': 20, 'groups': 5} for f in builder.FAMILIES},
               'traces': [{'cell_id': 'fixture', 'selected': 'The number is 2.', 'interpretation': 'Fixture only.'}],
               'cost_rows': [{'condition': c, 'components': [{'component': 'experimental_candidate_budget', 'value': 1},
                                                            {'component': 'recurring_helper', 'value': 0}]} for c in builder.METHODS]}
    report = {'signed_comparisons': {'validation': section}, 'activity': {'heldout_items': 268}}
    report_path = root/'report.json'; report_path.write_text(json.dumps(report))
    source = {'schema_version': 1, 'status': 'DRAFT - DEVELOPMENT / VALIDATION ONLY', 'test_quality_access': False,
              'cohorts': ['development', 'validation'], 'report_sha256': builder.digest(report_path),
              'validation_artifact_sha256': 'a'*64, 'source_updated_utc': 'fixture timestamp', 'slides': [
                  {'kind': 'question', 'title': 'Question', 'subtitle': 'Fixture only', 'notes': 'Fixture note'},
                  {'kind': 'validation', 'title': 'Validation', 'conditions': list(builder.METHODS), 'notes': 'Fixture note'},
                  {'kind': 'example', 'title': 'Example', 'cell_id': 'fixture', 'excerpts': ['The number is 2.'], 'interpretation': 'Fixture only.', 'notes': 'Fixture note'},
                  {'kind': 'cost', 'title': 'Cost', 'conditions': list(builder.METHODS), 'notes': 'Fixture note'}]}
    source_path = root/'source.json'; source_path.write_text(json.dumps(source))
    return source_path, report_path, source


class SlideTests(unittest.TestCase):
    def test_report_identity_and_test_quality_reject_before_output(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); source_path, report_path, source = fixture(root)
            source['test_quality_access'] = True; source_path.write_text(json.dumps(source))
            with self.assertRaisesRegex(ValueError, 'cannot consume TEST'): builder.load_inputs(source_path, report_path)
            source['test_quality_access'] = False; source['report_sha256'] = '0'*64; source_path.write_text(json.dumps(source))
            with self.assertRaisesRegex(ValueError, 'SHA mismatch'): builder.load_inputs(source_path, report_path)

    def test_changed_excerpt_rejected_and_metric_objects_preserved(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); source_path, report_path, source = fixture(root)
            source['slides'][2]['excerpts'] = ['The number is 3.']; source_path.write_text(json.dumps(source))
            with self.assertRaisesRegex(ValueError, 'exact retained'): builder.load_inputs(source_path, report_path)
            source['slides'][2]['excerpts'] = ['The number is 2.']; source_path.write_text(json.dumps(source))
            loaded = builder.load_inputs(source_path, report_path); pages, refs = builder.content(*loaded)
            self.assertEqual(len(refs), 16)
            for row in refs: self.assertEqual(row['metric'], loaded[3][row['family'], row['condition']])
            self.assertEqual(len(pages), 4)

    def test_pptx_contains_editable_objects_and_notes_and_refuses_overwrite(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        with tempfile.TemporaryDirectory() as d:
            root = Path(d); source_path, report_path, _ = fixture(root); output = root/'new'
            receipt = builder.build(source_path, report_path, output)
            deck = Presentation(output/'draft.pptx')
            self.assertEqual(len(deck.slides), 4)
            self.assertFalse(receipt['native_pptx_rendering_validated'])
            self.assertTrue(all(s.notes_slide.notes_text_frame.text for s in deck.slides))
            self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for s in deck.slides for shape in s.shapes))
            self.assertGreater(sum(len(s.shapes) for s in deck.slides), 150)
            with self.assertRaisesRegex(ValueError, 'immutable'): builder.build(source_path, report_path, output)


if __name__ == '__main__': unittest.main()
