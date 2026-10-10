"""Build a source-bound editable DRAFT PPTX and same-content HTML mirror.

Rendering only. No evaluation, inference or final TEST outcomes are accessed.
The HTML/PDF mirror is not a native PowerPoint rendering parity claim.
"""
import argparse
from hashlib import sha256
import html
from importlib.metadata import version
import json
import math
from pathlib import Path
import subprocess
import textwrap

W, H = 1600, 900
INK, MUTED, PAPER, TEAL, BLUE, LINE = '15304A', '5B6D7E', 'F2F6FA', '087F83', '246AB3', 'D6E1EC'
METHODS = {
    'raw_n1-n1': ('Original readout', '8495A7'),
    'best_n_exact-n16': ('Best of 16', BLUE),
    'best_n_exact-n32': ('Best of 32', '6194CB'),
    'best_n_specificity_mean-n32': ('Compare other activations', TEAL),
}
FAMILIES = {
    'hallucination': ('Transcript contradictions', 'Lower is better'),
    'multihop_mt': ('Multi-hop facts', 'Higher is better'),
    'chain_intermediates': ('Arithmetic steps', 'Higher is better'),
    'role_bound_association': ('Who did what', 'Higher is better'),
}


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def load_inputs(source_path, report_path):
    source = json.loads(Path(source_path).read_bytes())
    if source.get('status') != 'DRAFT - DEVELOPMENT / VALIDATION ONLY' or source.get('test_quality_access') is not False:
        raise ValueError('This style proof is explicitly draft and cannot consume TEST quality')
    if source.get('cohorts') != ['development', 'validation'] or source.get('schema_version') != 1:
        raise ValueError('Unknown slide-source cohort/schema')
    if digest(report_path) != source['report_sha256']:
        raise ValueError('Slide input report SHA mismatch')
    report = json.loads(Path(report_path).read_bytes())
    section = report['signed_comparisons']['validation']
    if section['review']['status'] != 'root_reviewed' or section['stage'] != 'validation':
        raise ValueError('Only reviewed validation data may enter this proof')
    if section['source_sha256'] != source['validation_artifact_sha256']:
        raise ValueError('Reviewed validation artifact mismatch')
    if len(source['slides']) != 4 or [s['kind'] for s in source['slides']] != ['question', 'validation', 'example', 'cost']:
        raise ValueError('Style proof has exactly four predefined slide types')
    for slide in source['slides']:
        if 'conditions' in slide and slide['conditions'] != list(METHODS):
            raise ValueError('Proof method list changed')
    metrics = {(r['family'], r['condition']): r['metric'] for r in section['cell_rows']}
    for family in FAMILIES:
        for condition in METHODS:
            m = metrics[family, condition]
            if not math.isfinite(m['value']) or not 0 <= m['value'] <= 1:
                raise ValueError('Invalid source metric')
            ci = m['ci95']
            if ci is not None and (len(ci) != 2 or not 0 <= ci[0] <= ci[1] <= 1):
                raise ValueError('Invalid source interval')
    example = source['slides'][2]
    retained = next(t for t in section['traces'] if t['cell_id'] == example['cell_id'])
    if any(q not in retained['selected'] for q in example['excerpts']) or example['interpretation'] != retained['interpretation']:
        raise ValueError('Example differs from exact retained reviewed text')
    for plot in source.get('existing_plots', []):
        if digest(plot['path']) != plot['sha256']:
            raise ValueError('Existing plot SHA mismatch')
    return source, report, section, metrics


def text(x, y, width, height, value, size=28, color=INK, bold=False, serif=False):
    return dict(kind='text', x=x, y=y, w=width, h=height, text=value,
                size=size, color=color, bold=bold, serif=serif)


def rect(x, y, width, height, color):
    return dict(kind='rect', x=x, y=y, w=width, h=height, color=color)


def line(x1, y1, x2, y2, color=LINE, width=2):
    return dict(kind='line', x=x1, y=y1, x2=x2, y2=y2, color=color, width=width)


def base(title, number, source):
    return [rect(0, 0, W, H, PAPER), text(60, 30, 1480, 35, 'ORACLE LENS / DRAFT STYLE PROOF / VALIDATION ONLY', 20, MUTED),
            text(60, 85, 1480, 95, title, 52, serif=True), line(60, 851, 1540, 851),
            text(60, 865, 1370, 25, 'Reviewed validation source ' + source['validation_artifact_sha256'][:12] + ' | No TEST quality results', 18, MUTED),
            text(1480, 865, 60, 25, str(number) + ' / 4', 18, MUTED)]


def content(source, report, section, metrics):
    pages = []
    metric_refs = []
    for index, slide in enumerate(source['slides'], 1):
        items = base(slide['title'], index, source)
        kind = slide['kind']
        if kind == 'question':
            items += [text(60, 195, 1390, 60, slide['subtitle'], 31, MUTED)]
            boxes = [(70, 'ONE SNAPSHOT', 'One layer and position'), (595, 'SHARED CANDIDATES', '1, 2, 4, 8, 16, 32'), (1120, 'SELECTED TEXT', 'Same frozen readers')]
            for x, heading, label in boxes:
                items += [rect(x, 320, 400, 175, 'FFFFFF'), text(x+25, 355, 350, 45, heading, 25, TEAL, True), text(x+25, 411, 350, 50, label, 27)]
            items += [text(480, 365, 100, 80, '→', 55, MUTED), text(1005, 365, 100, 80, '→', 55, MUTED),
                text(70, 565, 1460, 100, 'No overall winner is established.', 43, TEAL, True),
                text(70, 660, 1400, 130, 'Validation supports two experimental choices.\nRecovering a name or number can still leave the explanation wrong.', 32),
                text(70, 802, 1450, 35, 'The model and readers stay fixed. There is no new training.', 24, MUTED)]
        elif kind == 'validation':
            for i, family in enumerate(FAMILIES):
                x, y = 60+(i%2)*770, 190+(i//2)*325
                items += [rect(x, y, 710, 302, 'FFFFFF'), text(x+18, y+12, 465, 35, FAMILIES[family][0], 27, bold=True),
                          text(x+487, y+17, 210, 32, FAMILIES[family][1], 22, MUTED)]
                rows = [metrics[family, c] for c in slide['conditions']]
                maximum = min(1, max(.05, math.ceil(max((m['ci95'] or [m['value']])[-1] for m in rows)/.05)*.05))
                for j, (condition, m) in enumerate(zip(slide['conditions'], rows)):
                    name, color = METHODS[condition]; yy = y+70+j*52
                    items += [text(x+18, yy-18, 227, 44, name, 21, color), line(x+251, yy, x+505, yy)]
                    at = lambda value: x+251+254*value/maximum
                    if m['ci95']:
                        lo, hi = m['ci95']; items += [line(at(lo), yy, at(hi), yy, color, 3), line(at(lo), yy-7, at(lo), yy+7, color), line(at(hi), yy-7, at(hi), yy+7, color)]
                    items += [dict(kind='ellipse', x=at(m['value'])-6, y=yy-6, w=12, h=12, color=color),
                              text(x+530, yy-20, 161, 30, f"{m['value']*100:.1f}%", 26, color, True),
                              text(x+530, yy+12, 161, 28, f"{m['numerator']:,}/{m['denominator']:,}", 16, MUTED)]
                    metric_refs.append({'family': family, 'condition': condition, 'metric': m})
                for t in (0, .5, 1):
                    items += [text(x+251+254*t-23, y+253, 60, 25, f'{maximum*t*100:g}%', 17, MUTED)]
                items += [text(x+18, y+277, 675, 25, '20 items | ' + str(section['family_counts'][family]['groups'])+' related groups | 95% intervals', 18, MUTED)]
            items += [text(60, 822, 1470, 28, 'Contradictions / specific readouts; recovery / intended sites. Different tasks are not averaged.', 21, MUTED)]
        elif kind == 'example':
            items += [text(60, 197, 1450, 45, 'Reviewed validation example: arithmetic steps, layer 32', 28, MUTED),
                      rect(60, 280, 435, 420, 'E1EDF0'), text(88, 310, 380, 90, 'The number is recovered', 35, TEAL, True),
                      text(190, 410, 230, 170, '2', 130, TEAL, True), text(90, 594, 350, 75, 'That is the measured hit.\nIt is not a correct explanation.', 27),
                      text(550, 278, 970, 48, 'Exact excerpts from the selected readout', 30, bold=True)]
            yy = 352
            for excerpt in slide['excerpts']:
                lines = textwrap.fill(excerpt, width=56)
                items += [rect(540, yy-12, 980, 130, 'FFFFFF'), text(565, yy, 930, 110, lines, 30)]
                yy += 155
            items += [text(550, 690, 965, 100, slide['interpretation'], 30, MUTED),
                      text(60, 806, 1450, 33, slide['cell_id']+' | Official verdict unchanged', 20, MUTED)]
        elif kind == 'cost':
            items += [text(65, 196, 720, 48, 'Required samples per new activation', 30, bold=True),
                      text(805, 196, 725, 48, 'Runtime helper calls', 30, bold=True)]
            costs = {r['condition']: {c['component']: c for c in r['components']} for r in section['cost_rows']}
            for i, condition in enumerate(slide['conditions']):
                name, color = METHODS[condition]; yy = 292+i*75; row = costs[condition]
                items += [text(65, yy, 560, 70, name, 29, color), text(642, yy, 110, 55, str(row['experimental_candidate_budget']['value']), 39, color, True),
                          text(805, yy, 90, 55, str(row['recurring_helper']['value']), 39, color, True), text(910, yy+5, 600, 50, 'The benchmark judge is separate.', 25, MUTED)]
            items += [line(65, 604, 1530, 604), text(65, 639, 670, 65, str(report['activity']['heldout_items'])+' held-out items: quality pending', 32, TEAL, True),
                      text(65, 723, 680, 100, 'Cold setup, warm work and cache replay\nwill be reported separately.', 27, MUTED),
                      text(805, 639, 690, 65, 'No serving-time claim from cached work.', 32, bold=True),
                      text(805, 723, 690, 100, 'New-input latency and GPU dollar cost\nare missing in this reviewed snapshot.', 27, MUTED)]
        notes = slide['notes']+'\n\nSource report SHA-256: '+source['report_sha256']+'\nReviewed validation SHA-256: '+source['validation_artifact_sha256']+'\nSnapshot: '+source['source_updated_utc']
        pages.append({'title': slide['title'], 'items': items, 'notes': notes})
    return pages, metric_refs


def write_pptx(pages, target):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt
    prs = Presentation(); prs.slide_width = Inches(13+1/3); prs.slide_height = Inches(7.5)
    unit = lambda n: Inches(n/120)
    for page in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for item in page['items']:
            kind = item['kind']; color = RGBColor.from_string(item['color'])
            if kind == 'text':
                shape = slide.shapes.add_textbox(unit(item['x']), unit(item['y']), unit(item['w']), unit(item['h']))
                tf = shape.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                for i, value in enumerate(item['text'].split('\n')):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = value
                    p.font.name = 'Georgia' if item['serif'] else 'Arial'; p.font.size = Pt(item['size']*.6)
                    p.font.color.rgb = color; p.font.bold = item['bold']; p.space_before = p.space_after = 0; p.line_spacing = 1.15
            elif kind == 'line':
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, unit(item['x']), unit(item['y']), unit(item['x2']), unit(item['y2']))
                shape.line.color.rgb = color; shape.line.width = Pt(item['width']*.6)
            else:
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if kind == 'ellipse' else MSO_SHAPE.RECTANGLE, unit(item['x']), unit(item['y']), unit(item['w']), unit(item['h']))
                shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()
        slide.notes_slide.notes_text_frame.text = page['notes']
    prs.core_properties.title = 'Oracle Lens - DRAFT validation style proof'
    prs.core_properties.subject = 'Development and validation only; final TEST results are pending'
    prs.save(target)


def write_html(pages, target):
    slides = []
    for page in pages:
        children = []
        for o in page['items']:
            color = '#'+o['color']; xy = f"left:{o['x']*.8}px;top:{o['y']*.8}px;"
            if o['kind'] == 'text':
                style = xy+f"width:{o['w']*.8}px;height:{o['h']*.8}px;font-size:{o['size']*.8}px;color:{color};font-weight:{700 if o['bold'] else 400};font-family:{'Georgia' if o['serif'] else 'Arial'};"
                children.append('<div class="text" style="'+style+'">'+html.escape(o['text'])+'</div>')
            elif o['kind'] == 'line':
                children.append(f'<svg class="shape" width="{W*.8}" height="{H*.8}" aria-hidden="true"><line x1="{o["x"]*.8}" y1="{o["y"]*.8}" x2="{o["x2"]*.8}" y2="{o["y2"]*.8}" stroke="{color}" stroke-width="{o["width"]*.8}"/></svg>')
            else:
                style = xy+f"width:{o['w']*.8}px;height:{o['h']*.8}px;background:{color};border-radius:{'50%' if o['kind']=='ellipse' else '0'}"
                children.append('<div class="shape" style="'+style+'"></div>')
        slides.append('<section class="sheet"><div class="slide">'+''.join(children)+'</div></section>')
    target.write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>DRAFT validation style proof</title><style>@page{size:13.333333in 7.5in;margin:0}*{box-sizing:border-box}html,body{margin:0;background:#D6E1EC}.sheet{width:1280px;height:720px;position:relative;break-after:page;overflow:hidden}.sheet:last-child{break-after:auto}.slide{position:absolute;width:1280px;height:720px}.text,.shape{position:absolute}.text{white-space:pre-wrap;line-height:1.15}.shape{left:0;top:0;pointer-events:none}</style>'+''.join(slides)+'</html>')


def build(source_path, report_path, output):
    source, report, section, metrics = load_inputs(source_path, report_path)
    output = Path(output)
    if output.exists():
        raise ValueError('Use a new immutable draft output directory')
    pages, metric_refs = content(source, report, section, metrics)
    output.mkdir(parents=True)
    (output/'source.json').write_bytes(Path(source_path).read_bytes())
    (output/'report-source.json').write_bytes(Path(report_path).read_bytes())
    (output/'dependencies.json').write_text(json.dumps({p: version(p) for p in ('python-pptx', 'Pillow', 'lxml', 'XlsxWriter', 'PyMuPDF')}, indent=2)+'\n')
    write_pptx(pages, output/'draft.pptx'); write_html(pages, output/'draft.html')
    (output/'speaker-notes.md').write_text('\n\n'.join(f"# {i}. {p['title']}\n\n{p['notes']}" for i, p in enumerate(pages, 1))+'\n')
    (output/'geometry.json').write_text(json.dumps(pages, indent=2, ensure_ascii=False)+'\n')
    receipt = {'status': source['status'], 'test_quality_access': False, 'source_sha256': digest(source_path), 'report_sha256': digest(report_path),
               'builder_sha256': digest(__file__), 'slides': len(pages), 'metric_refs': metric_refs,
               'pdf_route': 'Same-geometry HTML mirror; not a native PPTX conversion', 'native_pptx_rendering_validated': False,
               'outputs': {p.name: digest(p) for p in output.iterdir() if p.is_file()}}
    (output/'build.json').write_text(json.dumps(receipt, indent=2)+'\n')
    return receipt


def render_pdf(output, chrome, node='node'):
    """Print the local HTML mirror and inspect actual PDF pages, without network."""
    output = Path(output).resolve(); receipt = json.loads((output/'build.json').read_bytes())
    if digest(output/'draft.html') != receipt['outputs']['draft.html']:
        raise ValueError('HTML mirror changed after its source-bound build')
    pdf = output/'draft.pdf'
    if pdf.exists():
        raise ValueError('Do not overwrite a rendered draft')
    driver = Path(__file__).with_name('render_slides.mjs').resolve()
    command = [str(node), str(driver), str(chrome), str(output/'draft.html'), str(pdf)]
    completed = subprocess.run(command, capture_output=True, timeout=45)
    (output/'pdf-driver-stderr.log').write_bytes(completed.stderr)
    if completed.returncode:
        raise RuntimeError('Local slide PDF conversion failed; diagnostic stderr retained')
    import pymupdf as fitz
    from PIL import Image
    document = fitz.open(pdf)
    if len(document) != receipt['slides']:
        raise ValueError('Unexpected printed PDF page count')
    previews = []
    for i, page in enumerate(document, 1):
        if abs(page.rect.width-960) > 1 or abs(page.rect.height-540) > 1:
            raise ValueError('Unexpected slide PDF dimensions')
        if 'ORACLE LENS / DRAFT STYLE PROOF / VALIDATION ONLY' not in page.get_text() or 'Reviewed validation source' not in page.get_text():
            raise ValueError('A slide header or source footer was clipped from the PDF')
        target = output/f'slide-{i:02d}.png'
        page.get_pixmap(matrix=fitz.Matrix(4/3, 4/3)).save(target)
        previews.append(target)
    document.close()
    sheet = Image.new('RGB', (1280, 720), '#D6E1EC')
    for i, preview in enumerate(previews):
        with Image.open(preview) as image:
            sheet.paste(image.resize((640, 360)), ((i%2)*640, (i//2)*360))
    sheet.save(output/'contact-sheet.png')
    rendered = {'pdf_route': receipt['pdf_route'], 'native_pptx_rendering_validated': False,
                'pages': len(previews), 'pdf_page_points': [960, 540], 'draft_pdf_sha256': digest(pdf),
                'build_sha256': digest(output/'build.json'), 'driver_sha256': digest(driver),
                'previews': {p.name: digest(p) for p in previews}}
    (output/'render.json').write_text(json.dumps(rendered, indent=2)+'\n')
    return rendered


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True); parser.add_argument('--report', type=Path, required=True); parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--chrome', type=Path, help='Optional local Chromium for HTML-mirror PDF and page previews')
    parser.add_argument('--node', default='node', help='Node >=22 for the optional bounded CDP renderer')
    args = parser.parse_args(); receipt = build(args.source, args.report, args.output)
    if args.chrome:
        render_pdf(args.output, args.chrome, args.node)
    print(json.dumps({k: receipt[k] for k in ('status', 'slides', 'pdf_route', 'native_pptx_rendering_validated')}))
